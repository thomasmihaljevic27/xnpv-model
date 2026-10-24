from pathlib import Path
from copy import deepcopy
from shutil import copy2
from docx import Document

src = Path('40_DOCS/Doc_3_Player_Pillar_II_2.docx')
archive = Path('90_ARCHIVE/2026-09-11/pricing_edit')
archive.mkdir(parents=True, exist_ok=True)
if not (archive / src.name).exists():
    copy2(src, archive / src.name)
d = Document(src)
original = [p.text for p in d.paragraphs]
r = Document('40_DOCS/Supervisor_Drafts/Data, Production, and Aging.docx')
replacements = {
5: "Of the 2,349 contracts, 529 pay the league minimum. A player cannot be paid less than that amount, even if his production would otherwise command a lower salary. The model handles this constraint with a censored regression, also called a Tobit model. For salaries above the minimum, it measures how closely the equation predicts observed pay. For minimum-salary contracts, it measures how likely the equation is to place the player at or below the minimum. This retains those observations without treating the salary floor as an unrestricted market price.",
7: "The sample begins in 2018 because earlier contract coverage is thin. The export contains 63, 223, and 452 contracts starting in 2015, 2016, and 2017, compared with roughly 700 a year afterward. The estimation uses the later window for its broader coverage.",
8: "The equation gives a player at zero WAR a price of about 1.325% of the cap, called the intercept. Each projected win adds 2.123% of the cap for a forward or 2.410% for a defenceman. These are estimated salary differences between positions, not differences in what a win contributes on the ice.",
9: "At a $95.5 million cap, a two-win forward is priced at about $5.32 million: $1.265 million at zero WAR plus two wins at $2.028 million each. A two-win defenceman is priced at about $5.87 million. Subtracting a $4 million cap hit leaves $1.32 million or $1.87 million in surplus for that season, before forecast uncertainty, exit risk, and discounting.",
12: "The main comparison used 1,781 contracts and measured production as average annual adjusted Game Value, converted to wins, over the three seasons after signing. Game Value contains no WAR input, so the outcome is constructed separately from the baseline. Each additional contract year was associated with -0.0135 wins per season, with a 95% interval from -0.0360 to +0.0089. Because this interval includes zero, the result did not justify increasing projected production for longer contracts.",
13: "For forwards alone, the association was +0.0273 wins per season per additional contract year, with a 95% interval from +0.0028 to +0.0518. This subgroup result has not been adopted in the price equation. Checks over two and four seasons also found no detectable association in the combined skater sample. These tests address future production; they do not establish what clubs pay for longer commitments, cost certainty, or flexibility.",
14: "A separate test found little evidence that RFA and UFA skaters required different price equations (F = 0.12; p = 0.887), so the model uses one equation for both groups. An exploratory test of star RFAs remains outside the model. Adopting that addition requires testing the star cutoff, refitting the current sample, and recalculating the affected valuations.",
15: "The equation holds the relationship between production and cap-adjusted pay constant across the sample. Measuring salaries as cap shares accounts for the cap's size but does not prove that this relationship is stable. A test comparing the pre-COVID growth years (2018-2019) with the flat-cap years (2020-2022) found no detectable change (p = 0.75).",
19: "For forecasts beyond three years, the model reuses the three-year spread, which can understate uncertainty on long contracts. The error estimates also come from positive-baseline players. Applying the same symmetric distribution to players whose forecasts have been moved to replacement level may misrepresent their upside and downside.",
21: "Historical forecasts for positive-baseline players missed production by an average absolute amount of 0.522 wins one year ahead, 0.735 wins two years ahead, and 0.803 wins three years ahead. Under the assumed normal distribution, multiplying these errors by 1.253 gives standard deviations of about 0.65, 0.92, and 1.01 wins. These describe the spread around the forecast. The model converts that spread into dollars using the position's price per win, then averages the possible prices after applying the minimum.",
27: "Actual renewals can involve arbitration or negotiated salaries and terms. Among 1,439 RFA-expiring deals with an observed next contract, the next cap hit equaled or exceeded the modeled offer in 48.3% of cases. The median difference was -$10,000 and the mean was +$690,000: the middle difference was close to zero, but more expensive renewals raised the average substantially. Where retention costs exceed the offer, the model overstates surplus. This comparison covers observed next contracts and their cap hits; it does not reconstruct individual negotiations.",
30: "Positive modeled surplus does not guarantee retention. For years in which keeping the player remains worthwhile, the model also weights surplus by the observed qualification rate for players of similar quality:",
35: "The salary equation otherwise makes many marginal players look worth retaining even though clubs let them go. Players without usable recent production remain in the bottom group to include likely departures. This expands the calibration sample from 1,231 to 2,270 decisions and raises the observed walk-away rate from 21.3% to 27.0%. The 100% star rate rests on only 22 decisions and treats that small sample's outcome as certain.",
37: "Qualification probabilities replace the exit probabilities used during the signed contract. The qualification chain starts at expiry without carrying forward the chance of departure before expiry; this remains unresolved. Control-year prices also apply the salary minimum directly, without the forecast-uncertainty adjustment used during the signed contract.",
39: "Only the contract active at the valuation date enters the contract-year calculation. A later extension cannot be taken from the historical record and included before it was agreed.",
40: "For each remaining season, the model multiplies projected production value by the probability of remaining in the NHL, subtracts the full cap cost, and discounts the surplus to the valuation date. The annual contribution is (S_k x V_k - C_k) / (1.03)^k, where k is the number of years after valuation, V_k is production value in dollars, C_k is cap cost, and S_k is the probability of remaining in the NHL.",
41: "The model sums these annual contributions and adds discounted RFA control-year surplus. In the valuation season, k is zero and survival is one, so neither survival weighting nor discounting reduces that season's value.",
43: "The model keeps the full cap cost even if the player leaves the NHL. It does not calculate contract-specific relief from retirement or other departures, so it overstates the cost where relief would apply.",
44: "The model assumes 3% annual cap growth from the ceiling known at valuation and discounts future surplus at the same rate. This expresses future dollars in the purchasing power of the valuation-date cap without using subsequently observed cap growth. The 3% rate is a planning assumption.\nThe model adds no discount for managerial impatience because the trade study is intended to investigate that behavior. A preference for near-term production would not establish impatience by itself: competitive position, uncertainty, and private information can also explain it.",
47: "Goalie performance is much less stable from year to year than skater performance. The model therefore places less weight on a goalie's recent results and more on league-average production. Goalies also have separate pricing and exit estimates.",
50: "Validation kept each goalie's records together when separating training and test samples. After adjusting shortened seasons, it supported placing 65% of the projection on the league mean.",
51: "The blended production forecast stays constant through the contract because the model has no fitted goalie aging curve. Validation supports the blend, not the choice to hold production constant.\nSeparate goalie exit probabilities reduce the contribution of future seasons as the chance of leaving the NHL accumulates. Forecasts retain the goalie's historical baseline even without a future performance record. A goalie with no usable history is left unpriced.",
52: "The goalie salary equation assigns 1.391% of the cap at zero WAR, plus 1.059% of the cap per goalie WAR. It uses an older ordinary least-squares regression, which fits salary to production without the skater regression's treatment of minimum-salary observations. Goalie valuations also use different salary-floor rules, and goalie control years lack the skater qualification-rate adjustment. These differences limit comparability even though both models express surplus in dollars."
}
# Rebuild using the reference package so theme fonts, styles and page settings match.
body = r._element.body
for child in list(body):
    if child.tag.endswith('}sectPr'):
        continue
    body.remove(child)
order = [i for i in range(53) if i != 4]
order.remove(19)
order.insert(order.index(21)+1,19)
for i in order:
    text = replacements.get(i,original[i]).replace('\u2019', "'").replace('\u2018', "'").replace('\u2014',' - ').replace('\u2013','-')
    style = 'Title' if i==0 else 'Heading 1' if i in [2,16,22,38,46] else 'Heading 2' if i in [24,28] else 'Normal'
    for part in text.split('\n'):
        p=r.add_paragraph(part,style)
        p.paragraph_format.widow_control=True
r.save(src)
check=Document(src)
assert len([p for p in check.paragraphs if p.style.name=='Heading 1'])==5
assert check.paragraphs[1].style.name=='Normal'
assert not any('\ufffd' in p.text for p in check.paragraphs)
print('Words:',sum(len(t.split()) for t in original),'->',sum(len(p.text.split()) for p in check.paragraphs))
print('Saved',src)
