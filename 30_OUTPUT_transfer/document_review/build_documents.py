from pathlib import Path
import json, math, csv, hashlib, re, copy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pdf2image import convert_from_path

ROOT=Path(__file__).resolve().parents[2]
WORK=Path(__file__).parent
OUT=ROOT/'40_DOCS'/'Model_Explanation_Review'
OUT.mkdir(exist_ok=True)
N=json.loads((ROOT/'30_OUTPUT/model_doc/numbers.json').read_text())
H=N['heavy']; P=N['prospects']; K=N['picks']
POP=Path(r'C:\Users\thoma\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin')
if not (POP/'pdftoppm.exe').exists(): POP=POP.parent.parent/'bin'
source=Document(r'C:\Users\thoma\Desktop\Valuing Players, Draft Picks, and Prospects - New Draft.docx')
source_tables=[[[c.text for c in row.cells] for row in t.rows] for t in source.tables]
def rows_csv(name):
 with (ROOT/'30_OUTPUT'/name).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

# Independent arithmetic checks against live output files, without rerunning fits.
checks=[]
def check(label, ok):
 assert ok,label
 checks.append(label)
line=json.loads((ROOT/'30_OUTPUT/xnpv1_price_line.json').read_text())['term_in']
for key in ['alpha','beta','beta_d_add','gamma_term']:
 check('price coefficient '+key,abs(line[key]-H['price_line']['locked']['term_in'][key])<1e-9)
for name,pl in H['players'].items():
 total=0
 for r in pl['seasons']:
  value=(r['p_play']*r['value_dollars']-r['cost_dollars']*(1-pl['retained_pct']))*r['discount']
  check(name+' season '+str(r['season_start']),abs(value-r['pv_dollars'])<.01)
  total+=value
 check(name+' total',abs(total-pl['summary']['npv_total'])<.01)
suz=next(r for r in rows_csv('prospect_values.csv') if r.get('player_name')=='Nick Suzuki' and r.get('trade_date','').startswith('2018-09-10'))
check('Suzuki current output',abs(float(suz['value_M'])-P['suzuki']['value'])<1e-8)
check('Suzuki components',abs(sum(P['suzuki']['parts'].values())-P['suzuki']['value'])<1e-9)
sis=next(r for r in rows_csv('pick_curve_players.csv') if r['player']=='Colton Sissons')
check('Sissons season sum',abs(sum(r['surplus'] for r in K['point']['seasons'])-float(sis['surplus'])/1e6)<1e-8)
check('Sissons control end',int(sis['window_end'])==2021)
check('Traded pick round-average calculation',abs(K['trade_pick']['star_chance']*P['star']['0']['b']-K['trade_pick']['value'])<1e-8)
check('Trade zero-sum',abs(N['trade']['MTL']+N['trade']['VGK'])<1e-9)
check('Montreal net',abs(N['trade']['sides']['MTL']['received_M']-N['trade']['sides']['MTL']['given_M']-N['trade']['MTL'])<1e-9)
current_fits=rows_csv('prospect_production_fits.csv')
(WORK/'verification.json').write_text(json.dumps({'passed':checks,'source_hashes':{f:hashlib.sha256((ROOT/f).read_bytes()).hexdigest() for f in ['20_CODE/skater_forecast.py','20_CODE/contract_npv.py','20_CODE/prospect_values.py','20_CODE/pick_curve.py','30_OUTPUT/model_doc/numbers.json','30_OUTPUT/xnpv1_price_line.json']},'scope':'Code-path review and independent output reconciliation; no model refit or full pipeline rerun.'},indent=2))

def plot(name,series,xlim,ylim,xticks,yticks,xlabel,ylabel,point=None):
 pdf=WORK/(name+'.pdf'); c=canvas.Canvas(str(pdf),pagesize=(600,310))
 l,b,w,h=64,55,514,205
 X=lambda x:l+(x-xlim[0])/(xlim[1]-xlim[0])*w
 Y=lambda y:b+(y-ylim[0])/(ylim[1]-ylim[0])*h
 c.setFont('Helvetica',10)
 for y in yticks:
  c.setStrokeColor(HexColor('#dddddd'));c.setLineWidth(.5);c.line(l,Y(y),l+w,Y(y))
  c.setFillColor(HexColor('#333333'));c.drawRightString(l-9,Y(y)-3,str(y))
 for x in xticks:c.drawCentredString(X(x),b-17,str(x))
 c.setStrokeColor(HexColor('#333333'));c.line(l,b,l+w,b);c.line(l,b,l,b+h)
 c.drawCentredString(l+w/2,17,xlabel);c.drawString(l,285,ylabel)
 colors=['#22577A','#B45309','#607C46']
 for i,(label,xs,ys) in enumerate(series):
  c.setStrokeColor(HexColor(colors[i]));c.setLineWidth(2)
  path=c.beginPath();path.moveTo(X(xs[0]),Y(ys[0]))
  for x,y in zip(xs[1:],ys[1:]):path.lineTo(X(x),Y(y))
  c.drawPath(path)
  c.line(l+i*174,272,l+18+i*174,272);c.setFillColor(HexColor('#222222'));c.drawString(l+23+i*174,269,label)
 if point:
  x,y,label=point;c.setFillColor(HexColor('#111111'));c.circle(X(x),Y(y),3.5,fill=1,stroke=0);c.drawString(X(x)+7,Y(y)+3,label)
 c.save();convert_from_path(str(pdf),dpi=160,poppler_path=str(POP))[0].save(WORK/(name+'.png'))

path=H['players']['Max Pacioretty']['forecast_path']
plot('aging',[('Pacioretty',[29+r['k'] for r in path],[r['rate_82'] for r in path])],(29,34),(0,1.8),list(range(29,35)),[0,.5,1,1.5],'Age on September 15','Forecast WAR per 82 games')
curve=K['shapes']['star probability (adopted)']['curve']
plot('picks',[('Adopted value',list(range(1,len(curve)+1)),curve)],(1,225),(0,15),[1,32,64,100,150,200],[0,5,10,15],'Draft position','Surplus ($M at the 2025-26 cap)')
co=P['suzuki']['coefs'];xs=[x/4 for x in range(101)]
series=[(f'Age {age}',xs,[co['p_star']*.169956+co['nhle82']*x+co['nhle82_sq']*x*x+co['prod_x_age18']*x*(age-18) for x in xs]) for age in [19,20,21]]
plot('prospect',series,(0,25),(-1,10),[0,5,10,15,20,25],[-1,0,2,4,6,8,10],'Production (NHLe points per 82 games)','D+2 prospect, drafted 13th: surplus ($M)',(P['suzuki']['inputs']['nhle82'],P['suzuki']['value'],'Suzuki'))

TABLES={}
def T(key,title,headers,rows,note,widths=None):TABLES[key]=(title,headers,rows,note,widths)
T('start_tests','Tests of the starting rate',['Choice / test','Result','Interpretation'],[
['Recent-season weights','50/30/20: 1.1708; 50/35/15: 1.1710','Starting-rate RMSE, WAR per 82. Small difference among nearby weights.'],
['Own-record share','55%: 1.1545; 65%: 1.1502; 75%: 1.1549','Starting-rate RMSE, WAR per 82; 65% lowest among these shares.'],
['Games weighting','0.8064 with; 0.8076 without; p = 0.026','Season-WAR RMSE, wins; games weighting improves the evaluated forecast.'],
['Minimum games','5 games: 0.8048; 10 games: 0.8064; p = 0.001','Season-WAR RMSE, wins. The five-game threshold scores better on evaluated rows.'],
['Rescaling missing seasons','Not scored','A design choice; missing history is not entered as zero performance.']],
'Sources: starting_level_simple_test.py and aging_investigations_abc.py (October 4–5, 2026); forecast_term_audit.py (October 21, 2026). Each comparison uses its own dated baseline. RMSE is root mean squared error: lower is better. The games-weighting and threshold tests cover 40,510 forecast rows; resampling is by career. Reported one-sided p-values are resampling comparison statistics.',[1.35,2.05,3.4])
T('comparables',"Pacioretty's five closest comparables at age 28",['Player','Similarity weight','Level (WAR per 82)'],[[r['name'],f"{r['weight']:.3f}",f"{r['level']:.3f}"] for r in H['players']['Max Pacioretty']['comparables']['top']], 'Source: the 2018-19 forecast reconstruction in model_doc/heavy.json. Five of 280 candidates; the full weighted group enters the forecast. Weights are not probabilities and need not sum to one.')
T('aging_tests','Departure assumptions and forecast error',['Departure value','Starting rate','Rate, 1–5 seasons','Season WAR'],[
['−0.50 WAR per 82','1.1494','1.3767','0.8139'],['−0.25 WAR per 82','1.1499','1.3670','0.8113'],['Replacement level (0)','1.1506','1.3598','0.8093'],['+0.25 WAR per 82','1.1515','1.3552','0.8080']],
'Source: aging_investigations_abc.py, October 5, 2026, 40,510 forecast rows across valuation seasons 2015-16 to 2021-22. The first two errors are RMSE in WAR per 82; the last is RMSE in season wins. These historical arms vary the departure value against that test’s baseline, before subsequent specification changes.')
T('participation','Valuation-season games share and chance of playing',['Input','Games share: coefficient (s.e.)','Playing log odds: coefficient (s.e.)'],[
['Recent games share','0.6246 (0.0132)','4.474 (0.153)'],['Recent season WAR','0.0376 (0.0024)','0.829 (0.057)'],['Seasons since debut','0.0049 (0.0009)','Not included'],['Age minus 27.22','−0.0101 (0.0008)','−0.274 (0.008)'],['Age squared after centring','0.0006 (0.0001)','Not included'],['Defenceman indicator','Not included','0.354 (0.066)'],['Under contract','Not included','2.088 (0.092)'],['Contract coverage unavailable','Not included','1.052 (0.075)'],['Intercept','0.1936 (0.0106)','−2.610 (0.103)']],
'Source: 2025-26 fits in model_doc/heavy.json. Games share: 11,844 player-seasons, 2,168 players, R² = 0.303. Playing: 16,421 player-seasons, 2,417 players, pseudo-R² = 0.379. Standard errors clustered by player; all displayed valuation-season coefficients have two-sided p < 0.001. Other horizons differ. The reported unpenalized playing coefficients differ from the production fit by at most 0.000008.')
T('participation_tests','Removing terms from the forecast',['Change to test baseline','Season-WAR RMSE','One-sided p, better'],[
['Baseline with all tested terms','0.8064','—'],['Remove defenceman from games share','0.8064','0.798'],['Remove age squared from playing probability','0.8062','0.013'],['Remove experience from playing probability','0.8056','<0.001'],['Remove all three (adopted)','0.8054','<0.001'],['Remove missing-contract-record indicator alone','0.8067','0.994']],
'Source: forecast_term_audit.py v1.2, October 21, 2026. RMSE in season wins, 40,510 forecast rows, valuation seasons 2015-16 through 2021-22. Each arm is compared with the test baseline, not with a previously modified arm. The missing-record removal worsens error (one-sided p = 0.006).')
pl=H['price_line']['fits']['wins + defence + term']['terms']
T('price','Current skater price regression',['Term','Cap-share coefficient','Dollar coefficient ($M)','s.e. ($M)','p'],[[lab,f"{r['b']:.6f}",f"{r['b']*95.5:.3f}",f"{r['se']*95.5:.3f}",'<0.001' if r['p']<.001 else f"{r['p']:.3f}"] for lab,r in pl.items()],
'Source: xnpv1_price_line.json and model_doc/heavy.json; October 21, 2026 lock. 2,336 contracts by 1,167 players; 521 censored at the minimum. Dollar coefficients use the $95.5M cap. Standard errors are clustered by player; p-values are two-sided. The additional defence coefficient adds to the forward coefficient.')
T('term_average','Which forecast best predicts cap hits?',['Forecast used','RMSE ($M)','Log likelihood','Lower RMSE than first season'],[[('First season' if r['version'].startswith('first') else 'Term average, conditional on playing' if r['version']=='term average (main)' else 'Term average, weighted by playing probability'),f"{float(r['rmse_M']):.4f}",f"{float(r['loglik']):.1f}",'—' if r['version'].startswith('first') else str(r['rmse_beats_first'])+' of 20'] for r in rows_csv('term_average_price_line.csv') if r['term']=='True'],
'Source: term_average_price_line.csv, October 21, 2026. Same 2,336 contracts. Five folds grouped by player, repeated 20 times. RMSE measures expected cap-hit prediction at the 2025-26 cap. Higher log likelihood is better in sample; lower RMSE is better out of sample.')
alt=source_tables[11]
T('price_alternatives','Alternative price specifications in the October 5 test',['Specification','Likelihood gain (p)','Held-out RMSE ($M)'],[r[:3] for r in alt[1:]],
'Source: price_line_spec_tests.py v1.0–v1.1, October 5, 2026; 2,296 contracts, grouped by player into five folds, 20 repeated splits. These results belong to the earlier coefficient fit and are not current-model validation. Likelihood p-values are likelihood-ratio tests. Values displayed as tied should not be interpreted as distinct.',[3.5,1.6,1.7])
cr=[]
for name,plr in H['players'].items():
 for r in plr['seasons']:
  y=int(r['season_start']);cr.append([name.split()[-1],f'{y}-{str(y+1)[2:]}',f"{r['projected_war']:.3f}",f"{r['p_play']:.3f}",f"{r['value_dollars']/1e6:.3f}",f"{r['cost_dollars']*(1-plr['retained_pct'])/1e6:.2f}",f"{r['pv_dollars']/1e6:.3f}"])
T('contracts','The contract seasons at the trade date',['Player','Season','WAR if playing','P(playing)','Price ($M)','Cost ($M)','PV ($M)'],cr,
'Source: current recorded forecast reconstruction, reconciled independently to each season’s PV. Dollars are at the valuation season’s purchasing power only after discounting; Price and Cost are nominal dollars for the row’s season. PV = (probability × price − cost) / 1.03ᵏ. Pacioretty totals −$1.958M; Tatar totals −$4.057M.',[1.05,.85,1,1,.95,.95,1])
T('control','Annual terminal-season weights for skaters',['Forecast WAR','Decisions','Qualify','Play if qualified','Product'],[
['3 or more','10','100.0%','100.0%','1.000'],['1 to <3','176','96.0%','100.0%','0.960'],['0 to <1','904','78.0%','88.0%','0.687'],['Below 0 / unavailable','1,180','65.4%','67.8%','0.443']],
'Source: rfa_terminal_value.py calibration recorded in model_doc/heavy.json. Contract expiries 2018-19 through 2024-25. The playing rate is measured among qualified players who appeared in the final contract season; the final-season playing probability is included separately. The top group has only ten decisions, so 1.000 is an empirical estimate, not certainty.')
T('goalies','Goaltender forecast and pricing checks',['Comparison','Recorded result','Metric / sample'],[
['Own-record share: 30%, 35%, 40%','7.3488 / 7.3294 / 7.3302','Season-WAR MSE; 928 seasons, 183 goalies'],['Linear age term versus no age term','7.2070 versus 7.3294; p = 0.022','Season-WAR MSE; grouped goalie holdouts'],['Age squared in playing probability','0.4850 with; 0.4845 without','Log loss; lower is better'],['Pricing forecast without / with age term','1.12 versus 1.27','Cap-hit RMSE ($M); repeated goalie holdouts']],
'Sources: goalie_review.py, goalie_play_spec.py, and goalie_line_check.py, October 21, 2026. MSE is mean squared error, in squared wins; it is not RMSE. The age-term improvement is 1.67% in MSE. Samples and targets differ across rows; their numbers are not directly comparable.')
T('sissons','Colton Sissons: realized surplus within modeled control',['Season','NHL games','Season WAR*','Value ($M)','Cost ($M)','Surplus ($M)'],[[f"{r['season']}-{str(r['season']+1)[2:]}",str(int(r['gp'])),f"{r['war82']:.3f}",f"{r['value']:.3f}",f"{r['cost']:.3f}",f"{r['surplus']:.3f}"] for r in K['point']['seasons']],
'Source: pick_curve_seasons.csv, reconciled to pick_curve_players.csv. *Season WAR is normalized for shortened league schedules; it is not WAR per 82 player games. All dollar columns are restated at the 2025-26 cap. Seasons without NHL games contribute zero. The control end-year is 2021, so 2020-21 is the final included season.')
T('pick_tests','Alternative shapes for draft-pick value',['Specification','#1 value ($M)','#32 value ($M)','Held-out RMSE ($M)'],[[name,f"{r['at']['1']:.2f}",f"{r['at']['32']:.2f}",f"{r['held']:.4f}"] for name,r in K['shapes'].items() if name in ['straight line in draft position','log of draft position','log of draft position and its square','star probability (adopted)','star probability and its square','star and NHL-player probabilities']],
'Source: model_doc/numbers.json, regenerated October 22, 2026. 2,316 players from eleven drafts. Each draft is predicted from a fit on the other ten; RMSE is in $M of career surplus. The star-probability input was estimated on overlapping historical drafts, so this is not a completely independent test of that input.',[3.5,1.05,1.05,1.2])
T('prospect_fade','Draft-position-only coefficient for remaining prospects',['Season','Prospect-seasons','Star value ($M)','s.e. ($M)'],[[('D+'+str(int(g)+1)) if g!='4+' else 'D+5 and later',f"{r['rows']:,}",f"{r['b']:.2f}",f"{r['se']:.2f}"] for g,r in P['star'].items()],
'Source: prospect_star_values.csv. All coefficients have p ≤ 0.001 in the recorded fit. The initial 2,316 players can appear in multiple seasons; 14,598 prospect-seasons are not 14,598 independent players. Standard errors are clustered by player. This table is the draft-position-only comparison, not the full scoring-and-age model.')
T('prospect_tests','Predictive comparisons for the prospect specification',['Specification','Held-out RMSE ($M)','Reported one-sided p'],[
['Draft position alone','2.2615','—'],['Draft position + linear production','2.2503','0.059 vs draft position alone'],['Draft position + linear and squared production','2.2158','0.033 vs draft position alone'],['Production terms; draft position only through D+3','2.2151','0.029 vs draft position alone'],['Above + production × age through D+4 (adopted)','2.1995','0.0015 vs preceding row'],['Linear production with age; omit squared term','2.2434','0.025, adopted better'],['Squared production with age; omit linear term','2.2012','0.337, adopted better']],
'Sources: prospect_nhle_look, prospect_age_look, and the production-form comparison in model_doc/numbers.json, using current recorded outcomes. 12,996 skater prospect-seasons from 2,083 players. Each draft is held out; comparison resampling is by draft. Different p-value baselines are named explicitly. The final restriction uses the code’s age basis, so it is not invariant to changing the age reference.',[3.5,1.2,2.1])
term_names={'p_star':'Draft probability','nhle82':'Production','nhle82_sq':'Production squared','prod_x_age18':'Production × (age − 18)'}
rr=[]
for g,fit in P['adopted'].items():
 for term,r in fit['terms'].items():rr.append([('D+'+str(int(g)+1)) if g!='4+' else 'D+5+',term_names[term],f"{r['b']:.4f}",f"{r['se']:.4f}",'<0.001' if r['p']<.001 else f"{r['p']:.3f}"])
T('prospect_coefficients','Full prospect regression, by season',['Season','Term','Coefficient','s.e.','p'],rr,
'Source: production fit reconstructed and guarded against prospect_production_fits.csv. Dependent variable: remaining surplus, $M at the 2025-26 cap. Standard errors clustered by player; two-sided p-values. The displayed age-minus-18 basis is an exact reparameterization of the code. D+1/D+2/D+3/D+4 rows: 2,083/2,030/1,906/1,715; D+5+ has 5,262 rows from 1,556 players. R² (uncentered): 0.292/0.220/0.158/0.106/0.035.',[.7,2.65,1.15,1.15,1.15])
T('suzuki',"Suzuki's value in the full prospect regression",['Term','Input','Coefficient','Contribution ($M)'],[[term_names[k],f"{P['suzuki']['inputs'][k]:.5f}",f"{P['suzuki']['coefs'][k]:.5f}",f'{v:.3f}'] for k,v in P['suzuki']['parts'].items()]+[['Total','','',f"{P['suzuki']['value']:.3f}"]],
'Source: prospect_values.csv and guarded reconstruction in model_doc/numbers.json. Contributions are calculated with unrounded coefficients and sum to $4.241747M. The inputs are observable player characteristics at the trade; fitted coefficients use the historical sample described in the text.',[2.75,1.35,1.35,1.35])
rr=[]
for a in N['trade']['assets']:
 ret=N['trade']['retained'].get(a['asset'],{}).get('pv_2025',0)
 rr.append([a['asset'],a['to'],f"{a['value']:.3f}",f'{-ret:.3f}' if ret else '—',f"{a['value']-ret:.3f}"])
T('trade','The assets in common-cap dollars',['Asset','Receiver','Value to receiver','Retained liability','Value before trade'],rr,
'Source: trade_asset_values.csv and trade_sides.csv, reconciled to recorded player and prospect values. All amounts $M at the 2025-26 cap. Retained liability belongs to the sending club; pre-trade value includes that liability. Rounded displayed cells can differ from unrounded totals by $0.001M.',[2, .65,1.4,1.4,1.35])

# Native Word math. Fractions and superscripts remain editable.
def mr(s):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
def frac(num,den):
 f=OxmlElement('m:f')
 for tag,val in [('m:num',num),('m:den',den)]:
  e=OxmlElement(tag);e.append(mr(val));f.append(e)
 return f
def eq(doc,key):
 if key=='participation':
  for text in ['g = a + b₁G + b₂E + b₃A + b₄A² + b₅W', 'p = 1 / (1 + exp(−z))']:
   p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
   m=OxmlElement('m:oMath');m.append(mr(text));p._p.append(m)
  return
 if key=='prospect':
  for label,formula in [('D+1 to D+3','aS + bP + cP² + eP(A − 18)'),('D+4','bP + cP² + eP(A − 18)'),('D+5 onward','bP + cP²')]:
   p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
   m=OxmlElement('m:oMath');m.append(mr(label+':  Value = '+formula));p._p.append(m)
  return
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_together=True
 m=OxmlElement('m:oMath')
 parts={
 'general':['Surplus = Σₖ ',('Vₖ − Cₖ','(1.03)ᵏ')],
 'performance':['wₖ = rₖ × gₖ;     E[WARₖ] = pₖ × wₖ'],
 'own':['Own rate = ',('0.5G₁r₁ + 0.3G₂r₂ + 0.2G₃r₃','0.5G₁ + 0.3G₂ + 0.2G₃')],
 'start':['Starting rate = 0.65 × own rate + 0.35 × L + Δ'],
 'price':['Cap share = α + (β + βᴅD)w + γT'],
 'contract':['Contract value = Σₖ ',('pₖVₖ − Cₖ','(1.03)ᵏ'),' + terminal value'],
 'terminal':['Terminal value = Σⱼ ',('Wⱼ(Vⱼ − Cⱼ)','(1.03)ʲ')],
 'historical':['Historical surplus = Σₛ ',('Vₛ − Cₛ','capₛ / cap*')],
 'pick':['Career surplus = θS + error;     pick value = θS'],
 'prospect':['D+1–D+3: aS + bP + cP² + eP(A − 18)\nD+4: bP + cP² + eP(A − 18)\nD+5 onward: bP + cP²']
 }[key]
 for item in parts:m.append(frac(*item) if isinstance(item,tuple) else mr(item))
 p._p.append(m)

def setup():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
 for st in ['Normal','Title','Heading 1','Heading 2','Caption']:
  s=d.styles[st];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
  s.paragraph_format.widow_control=True
 for s in d.styles:
  for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.12;d.styles['Normal'].paragraph_format.space_after=Pt(7)
 for st,sz in [('Title',21),('Heading 1',15),('Heading 2',12)]:
  d.styles[st].font.size=Pt(sz);d.styles[st].font.bold=True
  d.styles[st].paragraph_format.space_before=Pt(14);d.styles[st].paragraph_format.space_after=Pt(7)
 d.styles['Caption'].font.size=Pt(10);d.styles['Caption'].font.italic=False;d.styles['Caption'].font.bold=False
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run=f.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.append(fld)
 d.core_properties.author='';d.core_properties.last_modified_by='';return d

def table(doc,title,headers,rows,note,widths=None):
 p=doc.add_paragraph(title,'Caption');p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
 tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
 widths=widths or [6.8/len(headers)]*len(headers)
 for i,w in enumerate(widths):tb.columns[i].width=Inches(w)
 for row_i,vals in enumerate([headers]+rows):
  cells=tb.rows[0].cells if row_i==0 else tb.add_row().cells
  for j,val in enumerate(vals):
   cell=cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
   p.paragraph_format.keep_with_next=(len(rows)<=10 or row_i==0 or row_i==len(rows) or rows[row_i-1][0]==rows[row_i][0])
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(val))>30 else WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(str(val).replace('0.000 (','<0.001 ('));r.font.size=Pt(9.5);r.bold=row_i==0
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2' if row_i==0 else 'FFFFFF');pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
  trpr=tb.rows[row_i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if row_i==0:trpr.append(OxmlElement('w:tblHeader'))
 # Bold ties at displayed precision for comparative error columns.
 for j,h in enumerate(headers):
  if 'RMSE' in h or h in ['Starting rate','Rate, 1–5 seasons','Season WAR']:
   candidates=[]
   for i,row in enumerate(rows):
    try:candidates.append((float(str(row[j]).replace('$','')),i+1))
    except ValueError:pass
   if candidates:
    best=min(v for v,_ in candidates)
    for v,i in candidates:
     if v==best:
      for r in tb.rows[i].cells[j].paragraphs[0].runs:r.bold=True
 p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(10);p.paragraph_format.keep_together=True
 for r in p.runs:r.font.size=Pt(9)

def build(markdown,path):
 doc=setup();nt=nf=0
 for block in markdown.strip().split('\n\n'):
  if block.startswith('@table '):
   nt+=1;key=block.split()[1];title,head,rows,note,widths=TABLES[key];table(doc,f'Table {nt}. {title}',head,rows,note,widths)
  elif block.startswith('@figure '):
   nf+=1;key=block.split()[1];p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(WORK/(key+'.png')),width=Inches(6.65))
   caption={'aging':'Pacioretty’s forecast at the 2018 trade. Only 2018-19 belongs to the valued contract; later ages show the performance projection, not additional acquired seasons. Source: recorded 2018-19 forecast.', 'picks':'The adopted pick curve. The steep early decline follows the star-probability input. Source: current pick fit, 2,316 selections; dollars at the 2025-26 cap.', 'prospect':'Production and age in the D+2 valuation for a prospect drafted 13th. The curves hold draft position fixed and change age. Suzuki is plotted at his exact age and production, so he need not lie exactly on the age-19 line. Source: current prospect coefficients.'}[key]
   doc.add_paragraph(f'Figure {nf}. '+caption,'Caption')
  elif block.startswith('@eq '):eq(doc,block.split()[1])
  elif block.startswith('### '):doc.add_heading(block[4:],2)
  elif block.startswith('## '):doc.add_heading(block[3:],1)
  elif block.startswith('# '):doc.add_paragraph(block[2:],'Title')
  else:doc.add_paragraph(block)
 doc.save(path)
 with ZipFile(path) as z:
  assert z.testzip() is None
  for member in z.namelist():
   if member.endswith('.xml'):etree.fromstring(z.read(member))
 print(path,'tables',nt,'figures',nf)
 return doc

if __name__=='__main__':
 text=(WORK/'rewrite.md').read_text(encoding='utf-8').replace('−$0.205 million','−$0.206 million')
 build(text,OUT/'Valuing Players, Draft Picks, and Prospects - Independent Rewrite.docx')
 if (WORK/'review.md').exists():build((WORK/'review.md').read_text(encoding='utf-8'),OUT/'Review of the New Draft.docx')
 print('Checks passed:',len(checks))
